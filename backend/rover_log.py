"""Persistent multi-Rover recording and Excel export."""
from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import sqlite3
import threading

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook

from auth import require_api_key

router = APIRouter(prefix="/api/recording", dependencies=[Depends(require_api_key)])
DB_PATH = Path(__file__).with_name("rover_readings.sqlite3")
lock = threading.RLock()
recording = False

FIELDS = [
    "received_utc", "device_id", "latitude", "longitude", "altitude", "height_ellipsoid_m",
    "speed_kmh", "heading", "fix_quality", "fix_status_text", "fix_type", "fix_type_text",
    "battery_percent", "satellites_used", "hdop", "pdop", "horizontal_accuracy_m",
    "vertical_accuracy_m", "accuracy_source", "carrier_solution", "carrier_solution_text",
    "differential_solution", "ntrip_connected", "ntrip_status_text", "rtcm_bytes_injected",
    "rtcm_frames_injected", "correction_age_s", "last_rtcm_utc", "last_position_utc",
    "distance_to_base_meters", "precision_ready", "precision_block_reason",
]

INTEGER_FIELDS = {
    "fix_quality", "fix_type", "battery_percent", "satellites_used", "carrier_solution",
    "differential_solution", "ntrip_connected", "precision_ready", "rtcm_bytes_injected", "rtcm_frames_injected",
}
REAL_FIELDS = {
    "latitude", "longitude", "altitude", "height_ellipsoid_m", "speed_kmh", "heading", "hdop",
    "pdop", "horizontal_accuracy_m", "vertical_accuracy_m", "correction_age_s", "distance_to_base_meters",
}
COLUMN_TYPES = {name: ("INTEGER" if name in INTEGER_FIELDS else "REAL" if name in REAL_FIELDS else "TEXT") for name in FIELDS}


@contextmanager
def database():
    db = sqlite3.connect(DB_PATH)
    try:
        with db:
            columns = ", ".join(f"{name} {COLUMN_TYPES[name]}" for name in FIELDS)
            db.execute(f"CREATE TABLE IF NOT EXISTS readings (id INTEGER PRIMARY KEY, {columns})")
            existing = {row[1] for row in db.execute("PRAGMA table_info(readings)")}
            for name in FIELDS:
                if name not in existing:
                    db.execute(f"ALTER TABLE readings ADD COLUMN {name} {COLUMN_TYPES[name]}")
            yield db
    finally:
        db.close()


def append_reading(reading: dict) -> None:
    with lock:
        if not recording:
            return
        row = dict(reading, received_utc=datetime.now(timezone.utc).isoformat())
        with database() as db:
            columns = ",".join(FIELDS)
            placeholders = ",".join("?" for _ in FIELDS)
            db.execute(
                f"INSERT INTO readings ({columns}) VALUES ({placeholders})",
                [row.get(key) for key in FIELDS],
            )


@router.get("")
def status():
    with lock, database() as db:
        count = db.execute("SELECT COUNT(*) FROM readings").fetchone()[0]
        selected = ",".join(["id"] + FIELDS)
        rows = db.execute(f"SELECT {selected} FROM readings ORDER BY id DESC LIMIT 100").fetchall()
        return {"recording": recording, "count": count, "rows": [dict(zip(["id"] + FIELDS, row)) for row in rows]}


@router.post("/start")
def start():
    global recording
    with lock:
        recording = True
    return {"recording": True}


@router.post("/stop")
def stop():
    global recording
    with lock:
        recording = False
    return {"recording": False}


@router.delete("/clear")
def clear():
    with lock, database() as db:
        db.execute("DELETE FROM readings")
    return {"status": "cleared"}


@router.get("/export")
def export():
    with database() as db:
        selected = ",".join(["id"] + FIELDS)
        rows = db.execute(f"SELECT {selected} FROM readings ORDER BY id").fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RTK Rover Readings"
    headers = ["Reading ID"] + FIELDS
    sheet.append(headers)
    for row in rows:
        safe = ["'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value for value in row]
        sheet.append(safe)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(28, max(14, len(str(column[0].value)) + 2))

    header_index = {name: idx + 1 for idx, name in enumerate(headers)}
    for row in sheet.iter_rows(min_row=2):
        for key in ("latitude", "longitude"):
            row[header_index[key] - 1].number_format = "0.000000000"
        for key in ("horizontal_accuracy_m", "vertical_accuracy_m", "correction_age_s"):
            row[header_index[key] - 1].number_format = "0.0000"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Prosper_RTK_Rover_Readings.xlsx"'},
    )
