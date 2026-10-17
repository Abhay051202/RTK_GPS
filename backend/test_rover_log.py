import asyncio
import importlib
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

import rover_log


class RecordingTest(unittest.TestCase):
    def test_record_stop_persist_export(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.sqlite3"
            rover_log.DB_PATH = path
            reading = dict(
                device_id="=unsafe",
                latitude=0.0,
                longitude=73.12345678,
                fix_quality=4,
                fix_status_text="RTK FIXED",
                satellites_used=28,
                hdop=0.7,
                horizontal_accuracy_m=0.018,
                vertical_accuracy_m=0.031,
                carrier_solution=2,
                carrier_solution_text="RTK FIXED",
                differential_solution=True,
                rtcm_bytes_injected=12345,
                correction_age_s=0.4,
            )
            rover_log.stop()
            rover_log.append_reading(reading)
            self.assertEqual(rover_log.status()["count"], 0)
            rover_log.start()
            rover_log.append_reading(reading)
            rover_log.append_reading(reading)
            rover_log.stop()
            rover_log.append_reading(reading)
            self.assertEqual(rover_log.status()["count"], 2)

            importlib.reload(rover_log)
            rover_log.DB_PATH = path
            self.assertFalse(rover_log.status()["recording"])
            self.assertEqual(rover_log.status()["count"], 2)

            async def collect():
                return b"".join([chunk async for chunk in rover_log.export().body_iterator])

            workbook = load_workbook(BytesIO(asyncio.run(collect())))
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 3)
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers["latitude"]).value, 0)
            self.assertEqual(sheet.cell(2, headers["longitude"]).value, 73.12345678)
            self.assertEqual(sheet.cell(2, headers["fix_quality"]).value, 4)
            self.assertEqual(sheet.cell(2, headers["satellites_used"]).value, 28)
            self.assertNotEqual(sheet.cell(2, headers["device_id"]).data_type, "f")

    def test_migrates_legacy_schema(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "legacy.sqlite3"
            import sqlite3
            db = sqlite3.connect(path)
            db.execute("CREATE TABLE readings (id INTEGER PRIMARY KEY, received_utc TEXT, device_id TEXT, latitude REAL, longitude REAL, altitude REAL, speed_kmh REAL, heading REAL, fix_quality INTEGER, fix_status_text TEXT, battery_percent INTEGER)")
            db.commit(); db.close()
            rover_log.DB_PATH = path
            with rover_log.database() as migrated:
                columns = {row[1] for row in migrated.execute("PRAGMA table_info(readings)")}
            self.assertIn("horizontal_accuracy_m", columns)
            self.assertIn("rtcm_bytes_injected", columns)
            self.assertIn("last_rtcm_utc", columns)


if __name__ == "__main__":
    unittest.main()
